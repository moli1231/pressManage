from datetime import datetime
from io import BytesIO

from openpyxl import Workbook, load_workbook

from app.extensions import db
from app.models.gauge import GaugeType, PressureGauge


HEADERS = [
    "压力表编号",
    "检测时间(天)",
    "厂商",
    "编号",
    "放置位置",
    "实体照片",
    "检测报告照片",
    "校准日期",
    "类型(electronic/mechanical)",
]


def import_gauges_from_excel(file_stream, operator_id):
    wb = load_workbook(file_stream)
    ws = wb.active
    errors = []
    success = 0

    for idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        try:
            gauge_code, inspection_days, manufacturer, serial_no, location, entity_photo, report_photo, calibration_date, gauge_type = row

            if not all([gauge_code, inspection_days, manufacturer, serial_no, location, calibration_date, gauge_type]):
                raise ValueError("必填字段缺失")

            if gauge_type not in {GaugeType.ELECTRONIC.value, GaugeType.MECHANICAL.value}:
                raise ValueError("类型无效")

            calibration_date = _parse_date(calibration_date)
            inspection_days = int(inspection_days)

            existing = PressureGauge.query.filter_by(gauge_code=str(gauge_code)).first()
            if existing:
                raise ValueError("压力表编号重复")

            gauge = PressureGauge(
                gauge_code=str(gauge_code),
                inspection_days=inspection_days,
                manufacturer=str(manufacturer),
                serial_no=str(serial_no),
                location=str(location),
                entity_photo_path=str(entity_photo) if entity_photo else None,
                report_photo_path=str(report_photo) if report_photo else None,
                calibration_date=calibration_date,
                gauge_type=str(gauge_type),
                created_by=operator_id,
            )
            db.session.add(gauge)
            success += 1
        except Exception as exc:
            errors.append({"row": idx, "error": str(exc)})

    db.session.commit()
    return {"success": success, "errors": errors}


def export_gauges_to_excel(gauges):
    wb = Workbook()
    ws = wb.active
    ws.title = "压力表"
    ws.append(HEADERS)

    for gauge in gauges:
        ws.append(
            [
                gauge.gauge_code,
                gauge.inspection_days,
                gauge.manufacturer,
                gauge.serial_no,
                gauge.location,
                gauge.entity_photo_path,
                gauge.report_photo_path,
                gauge.calibration_date.isoformat() if gauge.calibration_date else "",
                gauge.gauge_type,
            ]
        )

    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return output


def _parse_date(value):
    if isinstance(value, datetime):
        return value.date()
    if hasattr(value, "isoformat"):
        return value
    return datetime.strptime(str(value), "%Y-%m-%d").date()
